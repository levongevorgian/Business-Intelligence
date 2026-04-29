import openpyxl
import pandas as pd
import pyodbc
import os
import configparser
import logging
from string import Template

# Setup logging
logging.basicConfig(
    filename='etl_process.log',
    level=logging.INFO,
    format='%(asctime)s %(levelname)s:%(message)s'
)

def build_connection_string(config_section):
    """
    Constructs a connection string from a configparser section.
    Ignores commented keys.
    """
    conn_str = ';'.join([f"{key}={value}" for key, value in config_section.items() if not key.startswith('#')])
    return conn_str

def main():
    # Read the configuration file
    config = configparser.ConfigParser()
    config_path = '../../infrastructure_initiation/sql_server_config.cfg'  # Adjust if needed

    if not os.path.exists(config_path):
        logging.error(f"Configuration file not found at path: {config_path}")
        return

    # Read the config file with environment variable substitution
    with open(config_path, 'r') as config_file:
        config_content = Template(config_file.read()).safe_substitute(os.environ)

    config.read_string(config_content)

    # Select the desired configuration section
    try:
        db_config = config['ORDER_DDS']
    except KeyError:
        logging.error("ORDER_DDS section not found in the configuration file.")
        return

    # Build the connection string
    connection_string = build_connection_string(db_config)
    logging.info("Connection string built successfully.")

    # Open database connection
    try:
        conn = pyodbc.connect(connection_string)
        cursor = conn.cursor()
        logging.info("Database connection established.")
    except pyodbc.Error as e:
        logging.error(f"Failed to connect to the database: {e}")
        return

    try:
        # Step 1: Populate Staging Tables
        logging.info("Populating Staging Tables...")
        excel_file_path = 'raw_data_source.xlsx'

        if not os.path.exists(excel_file_path):
            raise FileNotFoundError(f"Excel file does not exist at path: {excel_file_path}")

        workbook = openpyxl.load_workbook(excel_file_path)

        # Mapping from sheet names to suffixes
        sheet_to_suffix_map = {
            'Categories': 'Categories',
            'Customers': 'Customers',
            'Employees': 'Employees',
            'OrderDetails': 'OrderDetails',
            'Orders': 'Orders',
            'Products': 'Products',
            'Regions': 'Regions',     # Map both 'Regions' and 'Region' to 'Regions'
            'Region': 'Regions',
            'Shippers': 'Shippers',
            'Suppliers': 'Suppliers',
            'Territories': 'Territories'
        }

        # Primary key mapping for Dim_SOR population
        primary_key_map = {
            'Categories_Staging': 'staging_raw_category_id',
            'Customers_Staging': 'staging_raw_customer_id',
            'Employees_Staging': 'staging_raw_employee_id',
            'OrderDetails_Staging': 'staging_raw_orderdetail_id',
            'Orders_Staging': 'staging_raw_order_id',
            'Products_Staging': 'staging_raw_product_id',
            'Regions_Staging': 'staging_raw_region_id',
            'Shippers_Staging': 'staging_raw_shipper_id',
            'Suppliers_Staging': 'staging_raw_supplier_id',
            'Territories_Staging': 'staging_raw_territory_id'
        }

        # List to keep track of staging tables and their primary keys
        staging_tables_info = []

        for sheet_name in workbook.sheetnames:
            # Get table_suffix from mapping
            table_suffix = sheet_to_suffix_map.get(sheet_name.strip())
            if not table_suffix:
                logging.error(f"No mapping found for sheet name '{sheet_name}'. Skipping.")
                continue

            # Format table name as dbo.<TableSuffix>_Staging
            table_name = f"dbo.{table_suffix}_Staging"

            # Load data into pandas DataFrame
            df = pd.read_excel(excel_file_path, sheet_name=sheet_name)

            # Ensure valid column names
            if df.empty or any(df.columns.isnull()):
                logging.warning(f"Skipping sheet '{sheet_name}' due to missing/invalid columns.")
                continue

            # Data Validation: Truncate columns if necessary
            MAX_LENGTHS = {
                'PhotoPath': 200,
            }
            def truncate_data(row):
                for column, max_length in MAX_LENGTHS.items():
                    if column in row and isinstance(row[column], str):
                        row[column] = row[column][:max_length]
                return row
            df = df.apply(truncate_data, axis=1)

            # Ensure that 'RegionID' exists in Regions_Staging
            if table_suffix == 'Regions':
                if 'RegionID' not in df.columns:
                    logging.error(f"'RegionID' column missing in sheet '{sheet_name}'. Skipping.")
                    continue

            # Insert data into staging table using bulk insert
            placeholders = ', '.join(['?' for _ in df.columns])
            col_list = ', '.join(df.columns)

            sql = f"INSERT INTO {table_name} ({col_list}) VALUES ({placeholders})"

            try:
                cursor.fast_executemany = True
                cursor.executemany(sql, df.values.tolist())
                logging.info(f"Data loaded into staging table '{table_name}'.")
            except pyodbc.Error as e:
                logging.error(f"Failed to insert data into {table_name}: {e}")
                continue

            staging_tables_info.append({
                'table_suffix': f"{table_suffix}_staging",
                'table_name': table_name,
                'primary_key': primary_key_map.get(f"{table_suffix}_Staging")
            })

        # Commit staging table data
        conn.commit()
        logging.info("Staging table data committed.")

        # Step 2: Populate Dim_SOR
        logging.info("Populating Dim_SOR...")
        for staging_info in staging_tables_info:
            table_suffix = staging_info['table_suffix']
            table_name = staging_info['table_name']
            primary_key = staging_info['primary_key']

            if not primary_key:
                logging.error(f"No primary key defined for staging table '{table_name}'. Skipping Dim_SOR population.")
                continue

            sql = f"""
            INSERT INTO dbo.Dim_SOR (StagingTableName, StagingRawID, LoadDateTime)
            SELECT
                '{table_suffix}' AS StagingTableName,
                {table_name}.{primary_key} AS StagingRawID,
                GETDATE() AS LoadDateTime
            FROM {table_name}
            WHERE NOT EXISTS (
                SELECT 1 FROM dbo.Dim_SOR
                WHERE StagingTableName = '{table_suffix}' AND StagingRawID = {table_name}.{primary_key}
            );
            """
            try:
                cursor.execute(sql)
                logging.info(f"Dim_SOR populated for staging table '{table_suffix}'.")
            except pyodbc.Error as e:
                logging.error(f"Failed to populate Dim_SOR for '{table_suffix}': {e}")
                continue

        # Commit Dim_SOR data
        conn.commit()
        logging.info("Dim_SOR data committed.")

        # Step 3: Populate Dimension Tables
        logging.info("Populating Dimension Tables...")

        # Adjusted dimension scripts to remove DimSORID_fk and insert only existing columns
        dimension_scripts = {
            "DimCustomers": f"""
                INSERT INTO dbo.DimCustomers (
                    CustomerID_nk, CompanyName, ContactName, ContactTitle,
                    [Address], City, Region, PostalCode, Country, Phone, Fax
                )
                SELECT DISTINCT
                    stg.CustomerID,
                    stg.CompanyName,
                    stg.ContactName,
                    stg.ContactTitle,
                    stg.[Address],
                    stg.City,
                    stg.Region,
                    stg.PostalCode,
                    stg.Country,
                    stg.Phone,
                    stg.Fax
                FROM dbo.Customers_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'customers_staging' AND sor.StagingRawID = stg.staging_raw_customer_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimCustomers WHERE DimCustomers.CustomerID_nk = stg.CustomerID
                );
            """,
            "DimProducts": f"""
                INSERT INTO dbo.DimProducts (
                    ProductID_nk, ProductName, Supplier_sk_fk, Category_sk_fk,
                    QuantityPerUnit, UnitPrice, UnitsInStock, UnitsOnOrder,
                    ReorderLevel, Discontinued
                )
                SELECT DISTINCT
                    stg.ProductID,
                    stg.ProductName,
                    stg.SupplierID,
                    stg.CategoryID,
                    stg.QuantityPerUnit,
                    stg.UnitPrice,
                    stg.UnitsInStock,
                    stg.UnitsOnOrder,
                    stg.ReorderLevel,
                    stg.Discontinued
                FROM dbo.Products_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'products_staging' AND sor.StagingRawID = stg.staging_raw_product_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimProducts WHERE DimProducts.ProductID_nk = stg.ProductID
                );
            """,
            "DimCategories": f"""
                INSERT INTO dbo.DimCategories (
                    CategoryID_nk, CategoryName, [Description]
                )
                SELECT DISTINCT
                    stg.CategoryID,
                    stg.CategoryName,
                    stg.[Description]
                FROM dbo.Categories_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'categories_staging' AND sor.StagingRawID = stg.staging_raw_category_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimCategories WHERE DimCategories.CategoryID_nk = stg.CategoryID
                );
            """,
            "DimEmployees": f"""
                INSERT INTO dbo.DimEmployees (
                    EmployeeID_nk, LastName, FirstName, Title, TitleOfCourtesy,
                    BirthDate, HireDate, [Address], City, Region, PostalCode,
                    Country, HomePhone, Extension, Notes, ReportsTo, PhotoPath
                )
                SELECT DISTINCT
                    stg.EmployeeID,
                    stg.LastName,
                    stg.FirstName,
                    stg.Title,
                    stg.TitleOfCourtesy,
                    stg.BirthDate,
                    stg.HireDate,
                    stg.[Address],
                    stg.City,
                    stg.Region,
                    stg.PostalCode,
                    stg.Country,
                    stg.HomePhone,
                    stg.Extension,
                    stg.Notes,
                    stg.ReportsTo,
                    stg.PhotoPath
                FROM dbo.Employees_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'employees_staging' AND sor.StagingRawID = stg.staging_raw_employee_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimEmployees WHERE DimEmployees.EmployeeID_nk = stg.EmployeeID
                );
            """,
            "DimRegion": f"""
                INSERT INTO dbo.DimRegion (
                    RegionID_nk, RegionDescription
                )
                SELECT DISTINCT
                    stg.RegionID,
                    stg.RegionDescription
                FROM dbo.Regions_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'regions_staging' AND sor.StagingRawID = stg.staging_raw_region_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimRegion WHERE DimRegion.RegionID_nk = stg.RegionID
                );
            """,
            "DimShippers": f"""
                INSERT INTO dbo.DimShippers (
                    ShipperID_nk, CompanyName, Phone
                )
                SELECT DISTINCT
                    stg.ShipperID,
                    stg.CompanyName,
                    stg.Phone
                FROM dbo.Shippers_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'shippers_staging' AND sor.StagingRawID = stg.staging_raw_shipper_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimShippers WHERE DimShippers.ShipperID_nk = stg.ShipperID
                );
            """,
            "DimSuppliers": f"""
                INSERT INTO dbo.DimSuppliers (
                    SupplierID_nk, CompanyName, ContactName, ContactTitle,
                    [Address], City, Region, PostalCode, Country,
                    Phone, Fax, HomePage
                )
                SELECT DISTINCT
                    stg.SupplierID,
                    stg.CompanyName,
                    stg.ContactName,
                    stg.ContactTitle,
                    stg.[Address],
                    stg.City,
                    stg.Region,
                    stg.PostalCode,
                    stg.Country,
                    stg.Phone,
                    stg.Fax,
                    stg.HomePage
                FROM dbo.Suppliers_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'suppliers_staging' AND sor.StagingRawID = stg.staging_raw_supplier_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimSuppliers WHERE DimSuppliers.SupplierID_nk = stg.SupplierID
                );
            """,
            "DimTerritories": f"""
                INSERT INTO dbo.DimTerritories (
                    TerritoryID_nk, TerritoryDescription, TerritoryCode, Region_sk_fk
                )
                SELECT DISTINCT
                    stg.TerritoryID,
                    stg.TerritoryDescription,
                    stg.TerritoryCode,
                    reg.RegionID_sk_pk
                FROM dbo.Territories_Staging stg
                JOIN dbo.Dim_SOR sor
                    ON sor.StagingTableName = 'territories_staging' AND sor.StagingRawID = stg.staging_raw_territory_id
                JOIN dbo.DimRegion reg
                    ON stg.RegionID = reg.RegionID_nk
                WHERE NOT EXISTS (
                    SELECT 1 FROM dbo.DimTerritories WHERE DimTerritories.TerritoryID_nk = stg.TerritoryID
                );
            """,
        }

        for dim_table, sql_script in dimension_scripts.items():
            try:
                cursor.execute(sql_script)
                logging.info(f"{dim_table} populated successfully.")
            except pyodbc.Error as e:
                logging.error(f"Failed to populate {dim_table}: {e}")
                continue

        # Commit dimension table data
        conn.commit()
        logging.info("All dimension tables populated and committed.")

        # Step 4: Populate Fact Tables
        logging.info("Populating Fact Tables...")

        # FactOrders population
        fact_orders_script = f"""
            MERGE dbo.FactOrders AS DST
            USING (
                SELECT
                    stg.OrderID AS OrderID_nk,
                    stg.OrderDate,
                    stg.RequiredDate,
                    stg.ShippedDate,
                    c.CustomersID_sk_pk    AS Customer_sk_fk,
                    e.EmployeeID_sk_pk     AS Employee_sk_fk,
                    shp.ShippersID_sk_pk   AS ShipVia,
                    ter.TerritoriesID_sk_pk AS Territory_sk_fk,
                    stg.Freight,
                    stg.ShipName,
                    stg.ShipAddress,
                    stg.ShipCity,
                    stg.ShipRegion,
                    stg.ShipPostalCode,
                    stg.ShipCountry
                FROM dbo.Orders_Staging AS stg
                LEFT JOIN dbo.DimCustomers    AS c   ON stg.CustomerID  = c.CustomerID_nk
                LEFT JOIN dbo.DimEmployees    AS e   ON stg.EmployeeID  = e.EmployeeID_nk
                LEFT JOIN dbo.DimShippers     AS shp ON stg.ShipVia     = shp.ShipperID_nk
                LEFT JOIN dbo.DimTerritories  AS ter ON stg.TerritoryID = ter.TerritoryID_nk
            ) AS SRC
            ON DST.OrderID_nk = SRC.OrderID_nk

            WHEN MATCHED AND (
                   ISNULL(DST.Customer_sk_fk,  0) <> ISNULL(SRC.Customer_sk_fk,  0)
                OR ISNULL(DST.Employee_sk_fk,  0) <> ISNULL(SRC.Employee_sk_fk,  0)
                OR ISNULL(DST.ShipVia,         0) <> ISNULL(SRC.ShipVia,         0)
                OR ISNULL(DST.Territory_sk_fk, 0) <> ISNULL(SRC.Territory_sk_fk, 0)
                OR ISNULL(DST.OrderDate,       '1900-01-01') <> ISNULL(SRC.OrderDate,       '1900-01-01')
                OR ISNULL(DST.RequiredDate,    '1900-01-01') <> ISNULL(SRC.RequiredDate,    '1900-01-01')
                OR ISNULL(DST.ShippedDate,     '1900-01-01') <> ISNULL(SRC.ShippedDate,     '1900-01-01')
                OR ISNULL(DST.Freight,         0)            <> ISNULL(SRC.Freight,         0)
                OR ISNULL(DST.ShipName,        '')           <> ISNULL(SRC.ShipName,        '')
                OR ISNULL(DST.ShipAddress,     '')           <> ISNULL(SRC.ShipAddress,     '')
                OR ISNULL(DST.ShipCity,        '')           <> ISNULL(SRC.ShipCity,        '')
                OR ISNULL(DST.ShipRegion,      '')           <> ISNULL(SRC.ShipRegion,      '')
                OR ISNULL(DST.ShipPostalCode,  '')           <> ISNULL(SRC.ShipPostalCode,  '')
                OR ISNULL(DST.ShipCountry,     '')           <> ISNULL(SRC.ShipCountry,     '')
            )
            THEN
                UPDATE SET
                   DST.Customer_sk_fk   = SRC.Customer_sk_fk,
                   DST.Employee_sk_fk   = SRC.Employee_sk_fk,
                   DST.ShipVia          = SRC.ShipVia,
                   DST.Territory_sk_fk  = SRC.Territory_sk_fk,
                   DST.OrderDate        = SRC.OrderDate,
                   DST.RequiredDate     = SRC.RequiredDate,
                   DST.ShippedDate      = SRC.ShippedDate,
                   DST.Freight          = SRC.Freight,
                   DST.ShipName         = SRC.ShipName,
                   DST.ShipAddress      = SRC.ShipAddress,
                   DST.ShipCity         = SRC.ShipCity,
                   DST.ShipRegion       = SRC.ShipRegion,
                   DST.ShipPostalCode   = SRC.ShipPostalCode,
                   DST.ShipCountry      = SRC.ShipCountry

            WHEN NOT MATCHED BY TARGET
            THEN INSERT (
                OrderID_nk,
                Customer_sk_fk,
                Employee_sk_fk,
                ShipVia,
                Territory_sk_fk,
                OrderDate,
                RequiredDate,
                ShippedDate,
                Freight,
                ShipName,
                ShipAddress,
                ShipCity,
                ShipRegion,
                ShipPostalCode,
                ShipCountry
            )
            VALUES (
                SRC.OrderID_nk,
                SRC.Customer_sk_fk,
                SRC.Employee_sk_fk,
                SRC.ShipVia,
                SRC.Territory_sk_fk,
                SRC.OrderDate,
                SRC.RequiredDate,
                SRC.ShippedDate,
                SRC.Freight,
                SRC.ShipName,
                SRC.ShipAddress,
                SRC.ShipCity,
                SRC.ShipRegion,
                SRC.ShipPostalCode,
                SRC.ShipCountry
            )

            WHEN NOT MATCHED BY SOURCE
            THEN DELETE;
        """
        try:
            cursor.execute(fact_orders_script)
            logging.info("FactOrders populated successfully.")
        except pyodbc.Error as e:
            logging.error(f"Failed to populate FactOrders: {e}")

        # For OrderDetails
        order_details_script = f"""
            INSERT INTO dbo.OrderDetails (
                Order_sk_fk, Product_sk_fk, UnitPrice, Quantity, Discount
            )
            SELECT DISTINCT
                fo.FactOrdersID_sk_pk,
                dp.ProductID_sk_pk,
                stg.UnitPrice,
                stg.Quantity,
                stg.Discount
            FROM dbo.OrderDetails_Staging stg
            JOIN dbo.FactOrders fo ON stg.OrderID = fo.OrderID_nk
            JOIN dbo.DimProducts dp ON stg.ProductID = dp.ProductID_nk
            WHERE NOT EXISTS (
                SELECT 1 FROM dbo.OrderDetails
                WHERE Order_sk_fk = fo.FactOrdersID_sk_pk AND Product_sk_fk = dp.ProductID_sk_pk
            );
        """
        try:
            cursor.execute(order_details_script)
            logging.info("OrderDetails populated successfully.")
        except pyodbc.Error as e:
            logging.error(f"Failed to populate OrderDetails: {e}")

        # Commit fact table data
        conn.commit()
        logging.info("All fact tables populated and committed successfully.")

        logging.info("ETL process completed successfully.")

    except FileNotFoundError as e:
        logging.error(f"Error: The specified Excel file does not exist: {e}")
    except Exception as e:
        logging.error(f"An unexpected error occurred: {e}")
    finally:
        # Close the database connection
        try:
            cursor.close()
            conn.close()
            logging.info("Database connection closed.")
        except:
            logging.warning("Database connection was already closed or could not be closed.")

if __name__ == "__main__":
    main()
